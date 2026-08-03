# Copyright 2026 OpenSynergy Indonesia
# Copyright 2026 PT. Simetri Sinergi Indonesia
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).

import base64
import csv
import datetime
import hashlib
import io
import json

import openpyxl
import xlrd

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

from odoo.addons.ssi_decorator import ssi_decorator


class DataImport(models.Model):  # pylint: disable=too-few-public-methods
    """
    Transactional document that uploads a third-party file and, using
    the recipe stored on its ``data_import_template``, splits it into
    ``data_import.data`` lines -- one per source row, stored as JSON.
    Resolve (still owned by this model) finds each line's Target
    Model record and builds a Preview; on Queue Done,
    ``_10_enqueue_data_apply_jobs`` fans out one queue job per
    Matched line to actually write it (``data_import.data._apply``).
    Done is only reached once every line has left Draft/Matched/
    Error/Stale -- see ``_check_data_apply_complete``, checked both
    right after Queue Done (when Apply enqueued no job) and by this
    module's ``base.automation`` once the Done queue job batch
    finishes.

    Lifecycle: draft -> confirm -> queue_done -> done
    Cancellation: queue_cancel -> cancel (never touches Target Model
    records already written by Apply -- there is no rollback)
    """

    _name = "data_import"
    _description = "Data Import"
    _inherit = [
        "mixin.transaction_queue_cancel",
        "mixin.transaction_queue_done",
        "mixin.transaction_confirm",
    ]

    # View auto-insert attributes
    _automatically_insert_view_element = True
    _automatically_insert_multiple_approval_page = True
    _automatically_insert_done_policy_fields = False
    _automatically_insert_done_button = False
    _automatically_insert_queue_done_button = False
    _automatically_insert_queue_cancel_button = False
    _queue_processing_create_page = True

    # Approval attributes
    _approval_from_state = "draft"
    _approval_to_state = "action_queue_done"
    _approval_state = "confirm"
    _after_approved_method = "action_queue_done"

    # Sequence attribute
    _create_sequence_state = "queue_done"

    _statusbar_visible_label = "draft,confirm,queue_done,done"

    # Cancel-through-queue wizard attribute -- without this override, the
    # cancel reason wizard inherited from mixin.transaction_cancel would
    # call action_cancel() directly and skip the queue entirely.
    _method_to_run_from_wizard = "action_queue_cancel"

    _policy_field_order = [
        "confirm_ok",
        "approve_ok",
        "reject_ok",
        "restart_approval_ok",
        "queue_cancel_ok",
        "cancel_ok",
        "restart_ok",
        "done_ok",
        "queue_done_ok",
        "manual_number_ok",
    ]
    _header_button_order = [
        "action_confirm",
        "action_approve_approval",
        "action_reject_approval",
        "%(ssi_transaction_cancel_mixin.base_select_cancel_reason_action)d",
        "action_restart",
    ]
    _state_filter_order = [
        "dom_draft",
        "dom_confirm",
        "dom_reject",
        "dom_queue_done",
        "dom_done",
        "dom_terminate",
        "dom_queue_cancel",
        "dom_cancel",
    ]

    date = fields.Date(
        string="Date",
        required=True,
        default=lambda self: fields.Date.today(),
        readonly=True,
        states={"draft": [("readonly", False)]},
        help="Date of this import document.",
    )
    template_id = fields.Many2one(
        string="Template",
        comodel_name="data_import_template",
        required=True,
        ondelete="restrict",
        readonly=True,
        states={"draft": [("readonly", False)]},
        help=(
            "Data Import Template whose recipe is used to parse "
            "Import File: file format, offsets, and the Target Model "
            "below."
        ),
    )
    model_id = fields.Many2one(
        string="Target Model",
        comodel_name="ir.model",
        related="template_id.model_id",
        store=True,
        compute_sudo=True,
        help="Target Model configured on Template, kept in sync for reference.",
    )
    model_name = fields.Char(
        string="Target Model Technical Name",
        related="template_id.model_name",
        store=True,
        compute_sudo=True,
        help=(
            "Technical name of the Target Model configured on "
            "Template, kept in sync for reference."
        ),
    )
    import_file = fields.Binary(
        string="Import File",
        readonly=True,
        states={"draft": [("readonly", False)]},
        help="The source file to parse, in a format matching Template's File Format.",
    )
    import_file_name = fields.Char(
        string="File Name",
        readonly=True,
        states={"draft": [("readonly", False)]},
        help="Original file name of the uploaded Import File.",
    )
    import_file_hash = fields.Char(
        string="File Hash",
        readonly=True,
        copy=False,
        help=(
            "SHA-256 hash of Import File's content, computed by Load "
            "Data and used to detect duplicate imports."
        ),
    )
    data_ids = fields.One2many(
        string="Data",
        comodel_name="data_import.data",
        inverse_name="import_id",
        readonly=True,
        help=(
            "Rows read from Import File by Load Data, one record per "
            "source row. Rebuilt (not appended to) every time Load "
            "Data runs."
        ),
    )
    num_of_data = fields.Integer(
        string="# Data",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Total number of Data lines.",
    )
    num_of_matched = fields.Integer(
        string="# Matched",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Matched.",
    )
    num_of_done = fields.Integer(
        string="# Done",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Done.",
    )
    num_of_error = fields.Integer(
        string="# Error",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Error.",
    )
    num_of_no_match = fields.Integer(
        string="# No Match",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state No Match.",
    )
    num_of_multi_match = fields.Integer(
        string="# Multi Match",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Multiple Matches.",
    )
    num_of_conflict = fields.Integer(
        string="# Conflict",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Conflict.",
    )
    num_of_stale = fields.Integer(
        string="# Stale",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Stale.",
    )
    num_of_ignored = fields.Integer(
        string="# Ignored",
        compute="_compute_num_of_data",
        compute_sudo=True,
        help="Number of Data lines currently in state Ignored.",
    )
    resolved_date = fields.Datetime(
        string="Resolved Date",
        readonly=True,
        copy=False,
        help="Date and time Resolve last completed on this document.",
    )

    @api.depends("data_ids.state")
    def _compute_num_of_data(self):
        """Count Data lines per state for the header badge counters.

        :return: nothing; assigns the nine ``num_of_*`` counter fields
        """
        for record in self:
            data = record.data_ids
            result_total = len(data)
            result_matched = len(data.filtered(lambda d: d.state == "matched"))
            result_done = len(data.filtered(lambda d: d.state == "done"))
            result_error = len(data.filtered(lambda d: d.state == "error"))
            result_no_match = len(data.filtered(lambda d: d.state == "no_match"))
            result_multi_match = len(data.filtered(lambda d: d.state == "multi_match"))
            result_conflict = len(data.filtered(lambda d: d.state == "conflict"))
            result_stale = len(data.filtered(lambda d: d.state == "stale"))
            result_ignored = len(data.filtered(lambda d: d.state == "ignored"))
            record.num_of_data = result_total
            record.num_of_matched = result_matched
            record.num_of_done = result_done
            record.num_of_error = result_error
            record.num_of_no_match = result_no_match
            record.num_of_multi_match = result_multi_match
            record.num_of_conflict = result_conflict
            record.num_of_stale = result_stale
            record.num_of_ignored = result_ignored

    @api.constrains("import_file_hash")
    def _check_duplicate_import_file_hash(self):
        """Reject an Import File whose hash matches another live document.

        :raises ValidationError: when another ``data_import`` record
            (in any state but ``cancel``) already carries the same
            ``import_file_hash``
        """
        for document in self.sudo():
            if not document._check_duplicate_import_file_hash_condition():
                error_message = """
Context: Uploading data import file
Document: %s
Problem: This file has already been imported by another document
Solution: Check the existing import or use a different file""" % (
                    document.name or str(document.id)
                )
                raise ValidationError(_(error_message))

    def _check_duplicate_import_file_hash_condition(self):
        """Return whether ``import_file_hash`` is unique among live docs.

        :return: ``True`` when no other non-cancelled ``data_import``
            document shares this record's ``import_file_hash``
        """
        self.ensure_one()
        if not self.import_file_hash:
            return True
        duplicate = self.search(
            [
                ("import_file_hash", "=", self.import_file_hash),
                ("id", "!=", self.id),
                ("state", "!=", "cancel"),
            ],
            limit=1,
        )
        return not duplicate

    def action_load_data(self):
        """Parse Import File and rebuild Data lines for each record.

        :raises UserError: when a record is not in state ``draft``
        """
        for record in self.sudo():
            record._check_load_data_state()
            record._load_data()

    def _check_load_data_state(self):
        """Ensure Load Data only runs while the document is Draft.

        :raises UserError: when the document is not in state ``draft``
        """
        self.ensure_one()
        if self.state != "draft":
            error_message = """
Context: Loading data import file
Document: %s
Problem: Load Data is only allowed while the document is Draft
Solution: Reset the document to Draft before loading data again""" % (
                self.name or str(self.id)
            )
            raise UserError(_(error_message))

    def _load_data(self):
        """Rebuild ``data_ids`` from ``import_file``, idempotently.

        Existing ``data_ids`` are removed first, so calling this
        (via ``action_load_data``) more than once never duplicates
        rows. Also refreshes ``import_file_hash`` from the file
        content's SHA-256 digest before parsing.

        :return: nothing
        """
        self.ensure_one()
        if self.data_ids:
            self.data_ids.unlink()

        if not self.import_file:
            return

        file_content = base64.b64decode(self.import_file)
        self.import_file_hash = hashlib.sha256(file_content).hexdigest()

        rows = self._read_rows(file_content)
        data_vals = []
        for sequence, row in enumerate(rows, start=1):
            data_vals.append(
                {
                    "import_id": self.id,
                    "sequence": sequence,
                    "data": json.dumps(row),
                }
            )
        if data_vals:
            self.env["data_import.data"].create(data_vals)

    def _read_rows(self, file_content):
        """Return normalized row dicts read from ``file_content``.

        Applies Template's ``offset_row``, ``offset_column``,
        ``skip_empty_lines`` and ``no_header`` identically regardless
        of ``template_id.file_format``. Every cell value is turned
        into a string by ``_normalize_cell_value`` before being
        stored.

        :param file_content: raw bytes of the uploaded file
        :return: list of dict, one per data row, keyed by header name
            (or 0-based column index when ``no_header`` is checked)
        """
        self.ensure_one()
        template = self.template_id

        if template.file_format == "xlsx":
            raw_rows = self._iter_raw_rows_xlsx(file_content)
        elif template.file_format == "xls":
            raw_rows = self._iter_raw_rows_xls(file_content)
        else:
            raw_rows = self._iter_raw_rows_csv(file_content)

        rows = []
        headers = None
        for index, raw_row in enumerate(raw_rows):
            if template.offset_column:
                raw_row = raw_row[template.offset_column :]
            if index < template.offset_row:
                continue
            if template.skip_empty_lines and not any(
                str(cell).strip() for cell in raw_row
            ):
                continue
            if not template.no_header and headers is None:
                headers = [str(cell) for cell in raw_row]
                continue
            row_dict = {}
            for position, cell in enumerate(raw_row):
                key = headers[position] if headers else str(position)
                row_dict[key] = self._normalize_cell_value(cell, template.date_format)
            rows.append(row_dict)
        return rows

    def _normalize_cell_value(self, value, date_format):
        """Normalize a single raw cell value to a string for JSON storage.

        Empty becomes an empty string, booleans become ``"1"``/
        ``"0"``, integer-valued floats drop their decimal suffix
        (``"123"``, never ``"123.0"``) and ``date``/``datetime``
        values are formatted with ``date_format``.

        :param value: raw cell value from csv/openpyxl/xlrd
        :param date_format: ``strftime()`` format from Template's
            ``date_format``
        :return: normalized string value
        """
        if value is None:
            return ""
        if isinstance(value, bool):
            return "1" if value else "0"
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value.strftime(date_format or "%Y-%m-%d")
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return str(value)
        if isinstance(value, int):
            return str(value)
        return str(value)

    def _iter_raw_rows_csv(self, file_content):
        """Yield each row (list of cell strings) of a CSV import file.

        :param file_content: raw bytes of the uploaded file
        :raises UserError: when the bytes cannot be decoded using
            Template's ``file_encoding``
        :return: generator of row lists
        """
        self.ensure_one()
        template = self.template_id
        encoding = template.file_encoding or "utf-8"
        try:
            content_str = file_content.decode(encoding)
        except (LookupError, UnicodeDecodeError) as error:
            error_message = """
Context: Loading data import file
Document: %s
Problem: The uploaded file could not be decoded using encoding '%s'
Solution: Check the file matches Template's Encoding, or fix Encoding on the Template""" % (
                self.name or str(self.id),
                encoding,
            )
            raise UserError(_(error_message)) from error
        delimiter = template._get_delimiter_character()
        quotechar = template.quotechar or '"'
        lines_io = io.StringIO(content_str)
        reader = csv.reader(lines_io, delimiter=delimiter, quotechar=quotechar)
        for row in reader:
            yield list(row)

    def _iter_raw_rows_xlsx(self, file_content):
        """Yield each row (list of cell values) of an XLSX import file.

        :param file_content: raw bytes of the uploaded file
        :raises UserError: when the bytes are not a valid ``.xlsx``
            file
        :return: generator of row lists
        """
        self.ensure_one()
        template = self.template_id
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(file_content), read_only=True, data_only=True
            )
        except Exception as error:  # pylint: disable=broad-except
            error_message = """
Context: Loading data import file
Document: %s
Problem: The uploaded file could not be read as a valid Excel (.xlsx) file
Solution: Check the file is not corrupted and matches Template's File Format""" % (
                self.name or str(self.id)
            )
            raise UserError(_(error_message)) from error
        sheet = self._get_xlsx_sheet(workbook, template)
        for row in sheet.iter_rows(values_only=True):
            yield list(row)

    def _get_xlsx_sheet(self, workbook, template):
        """Return the worksheet selected by Template's Sheet Selector.

        :param workbook: an ``openpyxl`` workbook already loaded
        :param template: the ``data_import_template`` in use
        :return: an ``openpyxl`` worksheet
        """
        if template.sheet_selector == "name" and template.sheet_name:
            return workbook[template.sheet_name]
        return workbook.active

    def _iter_raw_rows_xls(self, file_content):
        """Yield each row (list of cell values) of an XLS import file.

        Date-typed cells (``xlrd`` type ``XL_CELL_DATE``) are
        converted to ``datetime.datetime`` using the workbook's date
        mode, so they are normalized the same way as ``.xlsx`` date
        cells.

        :param file_content: raw bytes of the uploaded file
        :raises UserError: when the bytes are not a valid ``.xls``
            file
        :return: generator of row lists
        """
        self.ensure_one()
        template = self.template_id
        try:
            workbook = xlrd.open_workbook(file_contents=file_content)
        except Exception as error:  # pylint: disable=broad-except
            error_message = """
Context: Loading data import file
Document: %s
Problem: The uploaded file could not be read as a valid Excel (.xls) file
Solution: Check the file is not corrupted and matches Template's File Format""" % (
                self.name or str(self.id)
            )
            raise UserError(_(error_message)) from error
        sheet = self._get_xls_sheet(workbook, template)
        for row_index in range(sheet.nrows):
            row = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                row.append(value)
            yield row

    def _get_xls_sheet(self, workbook, template):
        """Return the worksheet selected by Template's Sheet Selector.

        :param workbook: an ``xlrd`` workbook already opened
        :param template: the ``data_import_template`` in use
        :return: an ``xlrd`` sheet
        """
        if template.sheet_selector == "name" and template.sheet_name:
            return workbook.sheet_by_name(template.sheet_name)
        return workbook.sheet_by_index(0)

    def action_resolve(self):
        """Read-only, idempotent (re-)resolution of every Data line.

        Re-links every ``data_ids`` line currently in one of Draft,
        No Match, Multiple Matches, Conflict, Stale or Error to its
        Target Model record, using ``template_id``'s Matcher/Action
        rules, then flags conflicting lines. Never writes to the
        Target Model, and running it again after a previous run
        reaches the same result. Lines already Done or Ignored are
        left untouched. Available as a manual button while the
        document is Draft; also run automatically by
        ``_10_run_resolve`` right before Confirm's policy is enforced.

        :raises UserError: when the document is not in state ``draft``
        :return: nothing
        """
        for record in self.sudo():
            record._check_resolve_state()
            record._resolve()

    def _check_resolve_state(self):
        """Ensure Resolve only runs while the document is Draft.

        :raises UserError: when the document is not in state ``draft``
        """
        self.ensure_one()
        if self.state != "draft":
            error_message = """
Context: Resolving data import
Document: %s
Problem: Resolve is only allowed while the document is Draft
Solution: Reset the document to Draft before running Resolve again""" % (
                self.name or str(self.id)
            )
            raise UserError(_(error_message))

    def _resolve(self):
        """Re-resolve every eligible Data line, then flag conflicts.

        Called both by ``action_resolve`` (after its Draft guard) and
        directly by ``_10_run_resolve`` (while the document is still
        Draft, mid-Confirm, where the guard would be redundant).

        :return: nothing
        """
        self.ensure_one()
        lines = self.data_ids.filtered(lambda d: d.state in d._resolvable_states)
        for line in lines:
            line._resolve(self.template_id)
        self._check_resolve_conflicts()
        self.resolved_date = fields.Datetime.now()

    def _check_resolve_conflicts(self):
        """Flag intra-document duplicates and cross-document overlaps.

        Two lines in *this* document (state Matched or Done) pointing
        at the same Target Model record mark the later one (by
        ``sequence``) Conflict, with ``conflict_data_id`` set to the
        earlier one. The first line for each distinct target is then
        also checked against *other* documents -- see
        ``data_import.data._check_cross_document_conflict``.

        :return: nothing
        """
        self.ensure_one()
        seen = {}
        candidates = self.data_ids.filtered(
            lambda d: d.state in ("matched", "done")
            and d.source_document_model_id
            and d.source_document_res_id
        ).sorted("sequence")
        for line in candidates:
            key = (
                line.source_document_model_id.id,
                line.source_document_res_id,
            )
            first = seen.get(key)
            if first:
                line.write({"state": "conflict", "conflict_data_id": first.id})
                continue
            seen[key] = line
            line._check_cross_document_conflict()

    @ssi_decorator.pre_confirm_check()
    def _10_run_resolve(self):
        """Re-run Resolve immediately before Confirm's policy check.

        Keeps the gap between what the user reviewed and what a later
        Apply step will write as small as possible.

        :return: nothing
        """
        self._resolve()

    @ssi_decorator.pre_confirm_check()
    def _20_check_no_conflict(self):
        """Block Confirm while any Data line is in state Conflict.

        No Match and Multiple Matches lines never block Confirm --
        Template's On No Match / On Multiple Matches already resolved
        them by leaving the row alone -- only Conflict does.

        :raises UserError: when ``data_ids`` still has a Conflict
            line after ``_10_run_resolve`` just re-ran
        """
        self.ensure_one()
        conflict_lines = self.data_ids.filtered(lambda d: d.state == "conflict")
        if conflict_lines:
            error_message = """
Context: Confirm data import
Document: %s
Problem: %d Data line(s) are still in Conflict
Solution: Review the Conflict lines in the Import Data tab and re-run Resolve""" % (
                self.name or str(self.id),
                len(conflict_lines),
            )
            raise UserError(_(error_message))

    @ssi_decorator.post_queue_done_action()
    def _10_enqueue_data_apply_jobs(self):
        """Fan out one queue job per Matched Data line, on Queue Done.

        Runs inside ``action_queue_done``, right after the document's
        Done queue job batch (``done_queue_job_batch_id``) is
        created. Only lines already Matched by Resolve carry a
        Preview to apply -- No Match, Multiple Matches, Conflict,
        Ignored and Error lines are left untouched here, since none
        of them has one to apply. There is deliberately no matching
        ``post_queue_cancel_action`` hook: Cancel never touches
        ``data_ids`` (see the module's Keputusan Desain).

        :return: nothing
        """
        self.ensure_one()
        lines = self.data_ids.filtered(lambda d: d.state == "matched")
        for line in lines:
            description = "Apply data import line ID %s" % line.id
            job = (
                line.with_context(job_batch=self.done_queue_job_batch_id)
                .with_delay(description=_(description))
                ._apply()
            )
            line.queue_job_id = job.db_record().id

    def _set_done_if_no_job(self):
        """Complete Done immediately when Apply enqueued no job.

        Overrides ``mixin.transaction_queue_done``'s unconditional
        version (which reaches Done as soon as there is no queue
        job at all): Done is only reached here when, in addition,
        no Data line is left Draft, Matched, Error or Stale -- lines
        Resolve could not settle (e.g. a Required matcher's Column
        was empty) must not silently let the document reach Done
        just because Apply had nothing to enqueue.

        :return: nothing
        """
        self.ensure_one()
        if not self.done_queue_job_ids and self._check_data_apply_complete():
            self.action_done()

    def _recompute_queue_done_result(self):
        """Recompute the Done batch, completing Done once settled.

        Overrides ``mixin.transaction_queue_done``'s version, called
        by ``action_recompute_queue_done_result`` -- itself invoked
        by this module's ``base.automation`` whenever
        ``done_queue_job_batch_state`` becomes Finished -- to also
        require every Data line to have left Draft/Matched/Error/
        Stale before Done is reached.

        :return: nothing
        """
        self.ensure_one()
        self.done_queue_job_batch_id.enqueue()
        if (
            self.done_queue_job_batch_state == "finished"
            and self._check_data_apply_complete()
        ):
            self.action_done()

    def _check_data_apply_complete(self):
        """Return whether every Data line has settled for Done.

        :return: ``True`` when ``data_ids`` has no line left in
            state Draft, Matched, Error or Stale
        """
        self.ensure_one()
        blocking_states = ("draft", "matched", "error", "stale")
        return not self.data_ids.filtered(lambda d: d.state in blocking_states)

    @ssi_decorator.insert_on_form_view()
    def _insert_form_element(self, view_arch):
        if self._automatically_insert_view_element:
            view_arch = self._reconfigure_statusbar_visible(view_arch)
        return view_arch

    @api.model
    def _get_policy_field(self):
        res = super()._get_policy_field()
        policy_field = [
            "confirm_ok",
            "approve_ok",
            "reject_ok",
            "restart_approval_ok",
            "queue_done_ok",
            "queue_cancel_ok",
            "done_ok",
            "cancel_ok",
            "restart_ok",
            "manual_number_ok",
        ]
        res += policy_field
        return res
